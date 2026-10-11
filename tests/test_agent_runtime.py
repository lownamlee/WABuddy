"""Exercise real storage/queue/export boundaries; model responses are explicit fixtures."""
from io import BytesIO
import json
from pathlib import Path
import time
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
from PIL import Image

from wabuddy.agent_runtime import AgentRuntime
from wabuddy.agent_provider import Providers
from wabuddy.chat_store import ChatStore
from wabuddy.media import MediaStore
from test_app import app_client, login


class Registry:
    def __init__(self, root):
        self.root=root
        self.services={}
        for slot in ('a','b'):
            chats=ChatStore(root/slot/'chats.sqlite3')
            chats.activate(slot+'@s.whatsapp.net',slot)
            self.services[slot]=SimpleNamespace(chats=chats,snapshot=lambda:{'status':'connected'})

    def get(self,slot):
        return self.services[slot]


@pytest.fixture
def runtime(tmp_path):
    return AgentRuntime(Registry(tmp_path))


def seed(runtime, text='I confirm 2 boxes at RM34 each. Pickup 6 pm.', slot='a'):
    chats=runtime.registry.get(slot).chats
    chats.message('60111111111@s.whatsapp.net','order-1',text,False,100,sender='60111111111@s.whatsapp.net')
    return chats


def start(runtime, action='export_orders', slot='a', request='request-test-0001'):
    run=runtime.submit(slot,action,request)
    job=runtime.claim(slot)
    runtime.complete(slot,job['id'],job['lease'],{'action':action,'explanation':'Owner requested '+action})
    return run['id']


def analysis(job, quantity=2):
    return {'orders':[{'anchor':'order-1','item':'Boxes','quantity':quantity,'unit_price':34,'currency':'MYR',
        'status':'confirmed','payment':'unknown','delivery':'Pickup 6 pm',
        'evidence':[{'message_id':'order-1','quote':job['payload']['messages'][0]['text']}],'missing':[]}],
        'findings':[],'limitations':[],'image_observations':{}}


def finish(runtime, slot='a', quantity=2):
    job=runtime.claim(slot)
    runtime.complete(slot,job['id'],job['lease'],analysis(job,quantity))
    verify=runtime.claim(slot)
    assert verify['kind']=='verify'
    return runtime.complete(slot,verify['id'],verify['lease'],{'approved':True,'reason':'Fixture reviewer checked all source values.'})


def test_prompt_real_store_to_verified_xlsx_and_account_isolation(runtime):
    seed(runtime)
    run_id=start(runtime)
    result=finish(runtime)
    assert result['status']=='completed'
    assert len(result['nodes'])==4
    artifact=result['artifact']
    wb=load_workbook(runtime.artifact('a',artifact['id']),data_only=True)
    assert wb['Orders']['H2'].value==68
    assert wb['Evidence']['C2'].value=='order-1'
    assert wb['Summary']['B2'].value==1
    with pytest.raises(ValueError,match='not found'):
        runtime.run('b',run_id)
    with pytest.raises(ValueError,match='not found'):
        runtime.artifact('b',artifact['id'])
    assert runtime.context('a')['incoming_count']==1


def test_watch_updates_same_file_after_edit_and_survives_restart(runtime):
    chats=seed(runtime)
    start(runtime,'watch_orders')
    first=finish(runtime)
    runtime=AgentRuntime(runtime.registry)
    runtime.tick('a')
    assert runtime.claim('a') is None
    chats.update_message('60111111111@s.whatsapp.net','order-1',text='Confirmed: 3 boxes at RM34 each. Pickup 6 pm.',edited=True)
    runtime.tick('a')
    second=finish(runtime,quantity=3)
    assert first['artifact']['id']==second['artifact']['id']
    assert second['artifact']['revision']==2
    wb=load_workbook(runtime.artifact('a',second['artifact']['id']),data_only=True)
    assert wb['Orders'].max_row==2
    assert wb['Orders']['H2'].value==102
    runtime.tick('a')
    assert runtime.claim('a') is None


def test_schema_evidence_rejection_does_not_finish_job(runtime):
    seed(runtime)
    start(runtime)
    job=runtime.claim('a')
    data=analysis(job)
    data['orders'][0]['evidence'][0]['quote']='Invented confirmation'
    with pytest.raises(ValueError,match='quote'):
        runtime.complete('a',job['id'],job['lease'],data)
    assert runtime.run('a',job['run'])['nodes'][1]['status']=='working'
    data=analysis(job)
    data['orders'][0]['payment']='verified'
    with pytest.raises(ValueError):
        runtime.complete('a',job['id'],job['lease'],data)


@pytest.mark.parametrize('old_status', ['needs_input', 'failed'])
def test_old_owner_task_does_not_stop_order_monitoring(runtime, old_status):
    chats=seed(runtime)
    start(runtime,'watch_orders')
    first=finish(runtime)
    _,ns=runtime.identity('a')
    runtime.store.put(ns,'run','old-owner-task',{'id':'old-owner-task','source':'owner',
        'status':old_status,'prompt':'Unrelated earlier task','created':1})
    chats.update_message('60111111111@s.whatsapp.net','order-1',
        text='Confirmed: 3 boxes at RM34 each. Pickup 6 pm.',edited=True)
    runtime.tick('a')
    second=finish(runtime,quantity=3)
    assert second['artifact']['id']==first['artifact']['id']
    assert second['artifact']['revision']==2


@pytest.mark.parametrize('old_status', ['needs_input', 'failed'])
def test_blocked_scheduled_task_does_not_loop_model_calls(runtime, old_status):
    seed(runtime)
    start(runtime,'watch_orders')
    finish(runtime)
    _,ns=runtime.identity('a')
    runtime.store.put(ns,'run','blocked-schedule',{'id':'blocked-schedule','source':'schedule',
        'status':old_status,'prompt':'Process changes','created':1})
    runtime.registry.get('a').chats.message('60111111111@s.whatsapp.net','new','Another order',False,110)
    runtime.tick('a')
    assert runtime.claim('a') is None


def test_independent_rejection_creates_repair_then_needs_input(runtime):
    seed(runtime)
    start(runtime)
    first=runtime.claim('a')
    runtime.complete('a',first['id'],first['lease'],analysis(first))
    verify=runtime.claim('a')
    runtime.complete('a',verify['id'],verify['lease'],{'approved':False,'reason':'Quantity was misread.'})
    repair=runtime.claim('a')
    assert repair['payload']['correction']=='Quantity was misread.'
    assert repair['payload']['depends']==[verify['id']]
    runtime.complete('a',repair['id'],repair['lease'],analysis(repair))
    verify=runtime.claim('a')
    result=runtime.complete('a',verify['id'],verify['lease'],{'approved':False,'reason':'Still incorrect.'})
    assert result['status']=='needs_input'
    assert runtime.context('a')['artifacts']==[]


def test_leases_reclaim_cancel_and_prompt_idempotency(runtime):
    seed(runtime)
    run=runtime.submit('a','export orders','idempotent-request-1')
    assert runtime.submit('a','export orders','idempotent-request-1')['id']==run['id']
    with pytest.raises(ValueError):
        runtime.submit('a','another prompt','idempotent-request-1')
    first=runtime.claim('a')
    assert runtime.claim('a') is None
    with runtime.store.db() as db:
        db.execute('UPDATE jobs SET expires=0 WHERE id=?',(first['id'],))
    second=runtime.claim('a')
    assert first['lease']!=second['lease']
    with pytest.raises(ValueError,match='lease'):
        runtime.complete('a',first['id'],first['lease'],{'action':'export_orders','explanation':'x'})
    runtime.control('a',run['id'],'cancel')
    with pytest.raises(ValueError):
        runtime.complete('a',second['id'],second['lease'],{'action':'export_orders','explanation':'x'})


def test_parallel_claims_are_distinct_and_job_completion_idempotent(runtime):
    chats=seed(runtime)
    chats.message('60222222222@s.whatsapp.net','other','Hello',False,101)
    start(runtime)
    one=runtime.claim('a')
    two=runtime.claim('a')
    assert one['id']!=two['id']
    result=runtime.complete('a',one['id'],one['lease'],{'orders':[],'findings':[]})
    replay=runtime.complete('a',one['id'],one['lease'],{'orders':[],'findings':[]})
    assert len(result['nodes'])==len(replay['nodes'])


def test_owner_goal_scope_reaches_analysis_and_verifier(runtime):
    seed(runtime)
    prompt='Export only SAMPLE-ORDER test transactions, never actual sales.'
    runtime.submit('a',prompt,'scope-preservation-test')
    job=runtime.claim('a')
    runtime.complete('a',job['id'],job['lease'],{'action':'export_orders','explanation':'Export the requested test scope.'})
    job=runtime.claim('a')
    assert job['payload']['goal']==prompt
    assert job['payload']['owner_policies']==[]
    runtime.complete('a',job['id'],job['lease'],analysis(job))
    verifier=runtime.claim('a')
    assert verifier['payload']['goal']==prompt


def test_explicit_conversation_scope_avoids_unbounded_fanout(runtime):
    chats=seed(runtime)
    for index in range(6):
        chats.message(f'6099999900{index}@s.whatsapp.net',f'other-{index}','An unrelated conversation.',False,100)
    runtime.submit('a','Export orders only from 60111111111','bounded-scope-test')
    planner=runtime.claim('a')
    assert len(planner['payload']['available_conversations'])==7
    runtime.complete('a',planner['id'],planner['lease'],{'action':'export_orders','explanation':'Use the specified customer only.',
        'conversation_ids':['60111111111@s.whatsapp.net']})
    analyst=runtime.claim('a')
    assert analyst['payload']['chat']=='60111111111@s.whatsapp.net'
    assert runtime.claim('a') is None


def test_planner_rejects_unknown_chat_id_and_corrects_without_partial_work(runtime):
    from wabuddy.agent_provider import APIWorkers
    seed(runtime)
    run=runtime.submit('a','Export this customer order','planner-correction-test')
    job=runtime.claim('a')
    assert job['result_schema']['properties']['conversation_ids']['items']['enum']==['60111111111@s.whatsapp.net']
    attempts=[]
    class Provider:
        def infer(self, runtime, slot, request):
            attempts.append(request)
            return {'action':'export_orders','explanation':'Use the selected conversation.',
                'conversation_ids':['60111111111' if len(attempts)==1 else '60111111111@s.whatsapp.net']}
    workers=APIWorkers(runtime,Provider())
    try:
        workers._execute('a',job)
    finally:
        workers.close()
    assert len(attempts)==2
    assert 'available_conversations' in attempts[1]['payload']['validation_feedback']
    assert runtime.run('a',run['id'])['status']=='running'
    analyst=runtime.claim('a')
    assert analyst['payload']['chat']=='60111111111@s.whatsapp.net'
    assert runtime.claim('a') is None


def test_empty_account_planner_schema_cannot_select_a_chat(runtime):
    runtime.submit('a','Please prepare an order report','empty-planner-test')
    job=runtime.claim('a')
    assert job['result_schema']['properties']['conversation_ids']['maxItems']==0


def test_owner_handoff_quoted_reply_learning_is_customer_scoped(runtime):
    chats=seed(runtime,'Can I collect at 9 pm?')
    start(runtime,'learn_answers')
    job=runtime.claim('a')
    runtime.complete('a',job['id'],job['lease'],{'orders':[],'findings':[{
        'message_id':'order-1','kind':'question','reason':'No approved hours available.',
        'evidence':{'message_id':'order-1','quote':'Can I collect at 9 pm?'}}]})
    verifier=runtime.claim('a')
    runtime.complete('a',verifier['id'],verifier['lease'],{'approved':True,'reason':'Unknown question.'})
    assert len(runtime.context('a')['attention'])==1
    chats.message('60111111111@s.whatsapp.net','answer-1','Yes, for your order today.',True,102)
    chats.update_message('60111111111@s.whatsapp.net','answer-1',reply={'id':'order-1'})
    runtime.tick('a')
    memory=runtime.context('a')['knowledge'][0]
    assert memory['scope']=='customer'
    assert memory['answer']=='Yes, for your order today.'
    assert runtime.context('b')['knowledge']==[]


def test_phishing_filter_is_local_and_never_sends_or_deletes(runtime):
    seed(runtime,'Send me your OTP to avoid account suspension.')
    start(runtime,'filter_phishing')
    job=runtime.claim('a')
    runtime.complete('a',job['id'],job['lease'],{'orders':[],'findings':[{
        'message_id':'order-1','kind':'phishing','reason':'Requests a secret OTP under threat.',
        'evidence':{'message_id':'order-1','quote':'Send me your OTP'}}]})
    verifier=runtime.claim('a')
    runtime.complete('a',verifier['id'],verifier['lease'],{'approved':True,'reason':'Supported by source.'})
    assert runtime.context('a')['attention'][0]['filtered']
    assert runtime.context('a')['message_count']==1


@pytest.mark.parametrize('change', ['paused', 'enabled_after_scan', 'different_customer'])
def test_order_export_does_not_resolve_unreviewed_questions(runtime, change):
    seed(runtime)
    _,ns=runtime.identity('a')
    runtime.store.put(ns,'attention','pending-question',{'id':'pending-question',
        'chat':'60111111111@s.whatsapp.net','message_id':'question-1','kind':'question','status':'open'})
    policy={'id':'learn_answers','goal':'Review unknown questions','trigger':'message_changed',
        'enabled':change=='different_customer','conversation_ids':['60222222222@s.whatsapp.net'] if change=='different_customer' else []}
    runtime.store.put(ns,'policy',policy['id'],policy)
    start(runtime)
    if change=='enabled_after_scan':
        runtime.set_schedule('a','learn_answers',True)
    finish(runtime)
    assert runtime.store.get(ns,'attention','pending-question')['status']=='open'


def test_question_review_can_resolve_answered_questions(runtime):
    seed(runtime)
    _,ns=runtime.identity('a')
    runtime.store.put(ns,'attention','pending-question',{'id':'pending-question',
        'chat':'60111111111@s.whatsapp.net','message_id':'question-1','kind':'question','status':'open'})
    start(runtime,'learn_answers')
    finish(runtime)
    assert runtime.store.get(ns,'attention','pending-question')['status']=='resolved'


def test_image_access_requires_account_and_lease_and_retains_unknown(runtime):
    chats=seed(runtime,'[Image]')
    image=BytesIO()
    Image.new('RGB',(20,20),'green').save(image,format='PNG')
    media=MediaStore(chats).put('image','receipt.png','image/png',image.getvalue())
    chats.update_message('60111111111@s.whatsapp.net','order-1',media=media)
    start(runtime)
    job=runtime.claim('a')
    with pytest.raises(ValueError,match='Inspect each image'):
        runtime.complete('a',job['id'],job['lease'],{'orders':[],'findings':[],
            'image_observations':{'order-1':'Uninspected content'}})
    assert runtime.image('a',job['id'],job['lease'],'order-1').startswith(b'\xff\xd8')
    with pytest.raises(ValueError):
        runtime.image('b',job['id'],job['lease'],'order-1')
    runtime.complete('a',job['id'],job['lease'],{'orders':[],'findings':[]})
    verify=runtime.claim('a')
    assert 'Attachment not interpreted' in verify['payload']['candidate']['limitations'][0]


def test_provider_keys_never_returned_and_vision_no_text_fallback(runtime):
    providers=Providers(runtime.store.root)
    view=providers.save({'enabled':True,'text':{'base_url':'https://example.com/v1','model':'text','api_key':'private-test-key'}})
    assert 'private-test-key' not in json.dumps(view)
    providers.save({'enabled':True,'text':{'base_url':'https://example.com/v1','model':'new','api_key':None}})
    assert providers.read()['text']['api_key']=='private-test-key'
    with pytest.raises(ValueError,match='vision'):
        providers.infer(runtime,'a',{'run':'vision-test','payload':{'messages':[{'media':{'kind':'image'}}]}})
    with pytest.raises(ValueError,match='HTTPS'):
        providers.save({'text':{'base_url':'http://external.example/v1','model':'x'}})


def test_no_silent_history_truncation(runtime):
    chats=seed(runtime)
    with chats.db() as db:
        db.executemany('INSERT INTO messages(account,chat,id,text,outgoing,ts) VALUES(?,?,?,?,?,?)',
            [('a@s.whatsapp.net','60111111111@s.whatsapp.net',str(i),'text',0,i) for i in range(2001)])
    run_id=start(runtime)
    assert runtime.run('a',run_id)['status']=='needs_input'
    assert runtime.claim('a') is None


def test_live_http_session_csrf_and_no_key_leak(app_client):
    client,service,_=app_client
    assert client.get('/api/agent/context').status_code==401
    headers=login(client)
    service.chats.activate('60111111111@s.whatsapp.net','Test')
    body={'prompt':'Please prepare an order report','request_id':'request-id-test-001'}
    assert client.post('/api/agent/goals',json=body).status_code==403
    result=client.post('/api/agent/goals',headers=headers,json=body)
    assert result.status_code==200
    assert result.json()['status']=='waiting_for_agent'
    assert client.get('/api/agent/sharing-qr').content.startswith(b'\x89PNG')
    settings={'enabled':False,'text':{'base_url':'https://example.com/v1','model':'test','api_key':'private-secret'}}
    result=client.put('/api/agent/providers',headers=headers,json=settings)
    assert result.status_code==200
    assert 'private-secret' not in result.text
    assert client.get('/.local/agent/providers.json').status_code==404


def test_daily_schedule_fires_once_and_can_pause(runtime):
    seed(runtime)
    run=runtime.submit('a','Every evening at 10 pm export orders','daily-report-001')
    job=runtime.claim('a')
    result=runtime.complete('a',job['id'],job['lease'],{'action':'export_orders','explanation':'Daily at 22:00',
        'schedule':{'time':'22:00','timezone':'Asia/Singapore'}})
    assert result['scheduled']
    _,ns=runtime.identity('a')
    policy=runtime.context('a')['schedules'][0]
    policy['next_run']=time.time()-1
    runtime.store.put(ns,'policy',policy['id'],policy)
    runtime.tick('a')
    result=finish(runtime)
    assert result['source']=='schedule'
    runtime.tick('a')
    assert runtime.claim('a') is None
    runtime.set_schedule('a',policy['id'],False)
    assert not runtime.context('a')['schedules'][0]['enabled']


def test_completion_and_dependent_jobs_are_atomic(runtime,monkeypatch):
    seed(runtime)
    start(runtime)
    job=runtime.claim('a')
    def fail(*args,**kwargs):
        raise RuntimeError('Simulated process fault before dependent task creation')
    monkeypatch.setattr(runtime.store,'enqueue',fail)
    with pytest.raises(RuntimeError):
        runtime.complete('a',job['id'],job['lease'],analysis(job))
    _,ns=runtime.identity('a')
    assert runtime.store.leased(ns,job['id'],job['lease'])['status']=='working'


def test_unknown_prices_and_formula_text_are_preserved(runtime):
    seed(runtime,'Please order =HYPERLINK("https://example.test"). Price unknown.')
    start(runtime)
    job=runtime.claim('a')
    value=analysis(job)
    value['orders'][0].update(item='=HYPERLINK("https://example.test")',unit_price=None,currency=None)
    runtime.complete('a',job['id'],job['lease'],value)
    verify=runtime.claim('a')
    result=runtime.complete('a',verify['id'],verify['lease'],{'approved':True,'reason':'Unknown values kept blank.'})
    wb=load_workbook(runtime.artifact('a',result['artifact']['id']))
    assert wb['Orders']['D2'].data_type=='s'
    assert wb['Orders']['F2'].value is None
    assert wb['Orders']['H2'].value is None
    assert result['artifact']['totals']=={}


def test_api_provider_routes_image_to_vision_model(runtime,monkeypatch):
    from wabuddy import agent_provider
    seed(runtime)
    providers=Providers(runtime.store.root)
    providers.save({'enabled':True,
        'text':{'base_url':'https://text.example/v1','model':'text-model','api_key':'text-key'},
        'vision':{'base_url':'https://vision.example/v1','model':'vision-model','api_key':'vision-key'}})
    captured=[]
    class Opener:
        def open(self,request,timeout):
            captured.append(request)
            return BytesIO(json.dumps({'choices':[{'finish_reason':'stop','message':{'content':'{"orders":[],"findings":[]}'}}]}).encode())
    monkeypatch.setattr(agent_provider.urllib.request,'build_opener',lambda *args:Opener())
    monkeypatch.setattr(runtime,'image',lambda *args:b'fake-image-bytes')
    job={'id':'job','run':'vision-route-test','lease':'lease','instructions':'JSON only','result_schema':{},
         'payload':{'messages':[{'id':'image-message','media':{'kind':'image'}}]}}
    result=providers.infer(runtime,'a',job)
    assert result['orders']==[]
    request=captured[0]
    body=json.loads(request.data)
    assert request.full_url=='https://vision.example/v1/chat/completions'
    assert body['model']=='vision-model'
    assert body['messages'][1]['content'][-1]['type']=='image_url'
    assert request.headers['Authorization']=='Bearer vision-key'


def test_live_mcp_stdio_lists_real_job_tools(tmp_path):
    import asyncio
    import os
    import sys
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    params=StdioServerParameters(command=sys.executable,
        args=['-m','wabuddy.live_mcp','--account','default'],
        env={**os.environ,'PYTHONPATH':str(Path(__file__).resolve().parents[1]/'backend')})
    async def inspect():
        async with stdio_client(params) as (read,write):
            async with ClientSession(read,write) as session:
                await session.initialize()
                result=await session.list_tools()
                names={tool.name for tool in result.tools}
                assert {'get_execution_engine','set_execution_engine','read_whatsapp','update_whatsapp',
                        'send_message','send_attachment','message_action','search_messages',
                        'get_artifact','get_control_schema'} <= names
                assert {'claim_job','complete_job','inspect_image','submit_goal'}<=names
    asyncio.run(asyncio.wait_for(inspect(),timeout=25))


def test_live_dashboard_activity_names_and_workbook_preview(runtime):
    chats=seed(runtime)
    chats.message('60111111111@s.whatsapp.net','recent-in','A question',False,time.time())
    chats.message('60111111111@s.whatsapp.net','recent-out','Owner answer',True,time.time())
    chats.message('60111111111@s.whatsapp.net','deleted','Removed',False,time.time())
    chats.update_message('60111111111@s.whatsapp.net','deleted',deleted=True)
    context=runtime.context('a')
    assert sum(row['incoming'] for row in context['activity'])==1
    assert sum(row['outgoing'] for row in context['activity'])==1
    assert context['activity_timezone']=='UTC'
    assert runtime.context('b')['activity']==[]
    start(runtime)
    finish(runtime)
    context=runtime.context('a')
    assert context['order_preview'][0]['item']=='Boxes'
    assert context['order_preview'][0]['artifact_id']==context['artifacts'][0]['id']
    assert runtime.context('b')['order_preview']==[]


def test_owner_clarification_retains_context_and_unblocks_superseded_run(runtime):
    first=runtime.submit('a','Prepare a report','clarification-start-001')
    job=runtime.claim('a')
    runtime.complete('a',job['id'],job['lease'],{'action':'clarify','explanation':'What time each day?'})
    followup=runtime.submit('a','At 10 pm','clarification-answer-001',first['id'])
    assert runtime.run('a',first['id'])['status']=='cancelled'
    assert runtime.submit('a','At 10 pm','clarification-answer-001',first['id'])['id']==followup['id']
    job=runtime.claim('a')
    assert 'Prepare a report' in job['payload']['goal']
    assert 'At 10 pm' in job['payload']['goal']
    with pytest.raises(ValueError):
        runtime.submit('b','At 10 pm','clarification-answer-002',first['id'])


def test_revised_schedule_only_pauses_old_policy_when_replacement_is_valid(runtime):
    _,ns=runtime.identity('a')
    runtime.store.put(ns,'policy','old',{'id':'old','trigger':'daily','enabled':True,'goal':'Old time','schedule':{'time':'22:00','timezone':'Asia/Singapore'},'next_run':time.time()+86400})
    run=runtime.submit('a','Move report to 9 pm','replace-schedule-001',replace_schedule='old')
    assert runtime.store.get(ns,'policy','old')['enabled'] is True
    job=runtime.claim('a')
    runtime.complete('a',job['id'],job['lease'],{'action':'export_orders','explanation':'Daily at nine.','schedule':{'time':'21:00','timezone':'Asia/Singapore'}})
    assert runtime.store.get(ns,'policy','old')['enabled'] is False
    policies=runtime.context('a')['schedules']
    assert sum(p['enabled'] for p in policies)==1
    assert next(p for p in policies if p['enabled'])['schedule']['time']=='21:00'
    assert runtime.submit('a','Move report to 9 pm','replace-schedule-001',replace_schedule='old')['id']==run['id']
