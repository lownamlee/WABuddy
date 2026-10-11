from io import BytesIO
import asyncio
import os
import json
import sys
from pathlib import Path
import pytest
from openpyxl import load_workbook
from wabuddy.workspace import WorkspaceStore, RevisionConflict, ReportData, build_report
from wabuddy.workspace_mcp import create_server, validate_plan, PlannedTask
from test_app import app_client, login


def sample():
    return {'version':1, 'orders':[
        {'id':'1','customer':'Example','item':'Box','quantity':2,'unitPrice':34,'status':'Confirmed','payment':'Verified','delivery':'Tomorrow 4 pm','date':'2026-10-11','source':'=HYPERLINK("https://example.com")','revision':1},
        {'id':'2','customer':'Inquiry','item':'Cake','quantity':1,'unitPrice':120,'status':'Inquiry','payment':'Claimed','delivery':'','date':'2026-10-11','source':'How much?','revision':1},
    ],'threads':[{'id':'t1','customer':'Inquiry','reason':'Awaiting confirmation','resolved':False}],
    'knowledge':[],'automations':[],'runs':[],'reports':[],'attachments':[]}


def test_account_isolation_revision_conflicts_and_restart(tmp_path):
    store=WorkspaceStore(tmp_path/'state.db')
    assert store.save('a',0,sample())['revision']==1
    assert store.get('b')['state'] is None
    with pytest.raises(RevisionConflict): store.save('a',0,sample())
    assert WorkspaceStore(tmp_path/'state.db').get('a')['state']==sample()


def test_verified_workbook_separates_inquiries_and_preserves_text(tmp_path):
    store=WorkspaceStore(tmp_path/'state.db');store.save('a',0,sample())
    result=store.report('a',1,'2026-10-11')
    workbook=load_workbook(BytesIO(result['content']),data_only=False)
    assert workbook.sheetnames==['Orders','Summary','Needs Attention']
    assert workbook['Summary']['B7'].value==68
    assert workbook['Summary']['B8'].value==68
    assert workbook['Orders']['K3'].data_type=='s'
    assert workbook['Orders']['K3'].value.startswith('=HYPERLINK')
    assert store.get('a')['state']['reports'][0]['id']==result['report']['id']
    with pytest.raises(RevisionConflict): store.report('a',1,'2026-10-11')
    report=ReportData.model_validate(result['report']);report.orders.append(report.orders[0])
    with pytest.raises(ValueError,match='Duplicate'): build_report(report)


def test_workspace_http_requires_session_csrf_and_revision(app_client):
    client,_,_=app_client
    assert client.get('/api/workspace').status_code==401
    headers=login(client)
    assert client.post('/api/workspace',json={'expected_revision':0,'state':sample()}).status_code==403
    assert client.post('/api/workspace',headers=headers,json={'expected_revision':0,'state':sample()}).status_code==200
    assert client.get('/api/workspace',headers=headers).json()['state']['orders'][0]['id']=='1'
    assert client.post('/api/workspace',headers=headers,json={'expected_revision':0,'state':sample()}).status_code==409


def test_dependency_graph_rejects_cycles_and_unknown_dependencies():
    def task(id,depends):return PlannedTask(id=id,title=id,role='Coordinator',kind='Tool',depends=depends)
    validate_plan([task('a',[]),task('b',['a'])])
    with pytest.raises(ValueError):validate_plan([task('a',['b']),task('b',['a'])])
    with pytest.raises(ValueError):validate_plan([task('a',['missing'])])


def test_actual_mcp_tool_discovery_and_report(tmp_path):
    store=WorkspaceStore(tmp_path/'workspace.sqlite3');store.save('test-account',0,sample())
    server=create_server(tmp_path,'test-account')
    async def inspect():
        tools=await server.list_tools()
        assert {t.name for t in tools}=={'get_workspace_context','get_capabilities','create_goal_plan','get_run','list_outcomes','generate_order_report'}
        await server.call_tool('generate_order_report',{'scope':'2026-10-11','expected_revision':1})
    asyncio.run(inspect())
    assert len(list((tmp_path/'artifacts').glob('*.xlsx')))==1
    assert len(store.get('test-account')['state']['reports'])==1


def test_mcp_stdio_transport(tmp_path):
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client
    params=StdioServerParameters(command=sys.executable,
        args=['-m','wabuddy.workspace_mcp','--data-dir',str(tmp_path),'--account','test-account'],
        env={**os.environ,'PYTHONPATH':str(Path(__file__).resolve().parents[1]/'backend')})
    async def connect():
        async with stdio_client(params) as (read,write):
            async with ClientSession(read,write) as session:
                await session.initialize()
                result=await session.list_tools()
                assert len(result.tools)==6
                value=await session.call_tool('get_capabilities',{})
                assert not value.isError
                assert json.loads(value.content[0].text)['live_sending'] is False
    asyncio.run(asyncio.wait_for(connect(),timeout=25))
