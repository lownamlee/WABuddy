"""Account-scoped UI workspace and verified sample reports. Never sends messages."""
from __future__ import annotations

from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import json
import sqlite3
from typing import Literal
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel, Field


class ReportOrder(BaseModel):
    id: str = Field(min_length=1, max_length=100)
    customer: str = Field(max_length=300)
    item: str = Field(min_length=1, max_length=500)
    quantity: int = Field(ge=1, le=10000)
    unitPrice: Decimal = Field(ge=0, le=1000000, decimal_places=2)
    status: Literal['Inquiry', 'Draft', 'Confirmed', 'Cancelled']
    payment: Literal['Unpaid', 'Claimed', 'Receipt review', 'Verified']
    delivery: str = Field(max_length=1000)
    source: str = Field(max_length=10000)
    date: str = Field(pattern=r'^\d{4}-\d{2}-\d{2}$')
    revision: int = Field(ge=1)


class ReportData(BaseModel):
    id: str = Field(max_length=100)
    date: str = Field(pattern=r'^\d{4}-\d{2}-\d{2}$')
    range: str = Field(max_length=100)
    orders: list[ReportOrder] = Field(max_length=2000)
    attention: list[str] = Field(max_length=2000)
    coverage: str = Field(max_length=3000)
    created: str = Field(max_length=100)


class RevisionConflict(ValueError):
    pass


class WorkspaceStore:
    def __init__(self, path: Path):
        self.path = path
        path.parent.mkdir(parents=True, exist_ok=True)
        with sqlite3.connect(path) as db:
            db.execute('CREATE TABLE IF NOT EXISTS workspace (account TEXT PRIMARY KEY, revision INTEGER NOT NULL, state TEXT NOT NULL)')

    def get(self, account: str) -> dict:
        with sqlite3.connect(self.path) as db:
            row = db.execute('SELECT revision, state FROM workspace WHERE account=?', (account,)).fetchone()
        return {'revision': row[0], 'state': json.loads(row[1])} if row else {'revision': 0, 'state': None}

    def save(self, account: str, expected_revision: int, state: dict) -> dict:
        if state.get('version') != 1 or any(not isinstance(state.get(key), list) for key in ('orders','threads','knowledge','automations','runs','reports','attachments')):
            raise ValueError('Invalid workspace document.')
        text = json.dumps(state, ensure_ascii=False, allow_nan=False)
        if len(text.encode('utf-8')) > 2_000_000:
            raise ValueError('Workspace exceeds the 2 MB preview limit.')
        with sqlite3.connect(self.path, timeout=10) as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT revision FROM workspace WHERE account=?', (account,)).fetchone()
            current = row[0] if row else 0
            if current != expected_revision:
                raise RevisionConflict('Workspace changed in another window or tool. Reload before saving.')
            revision = current + 1
            db.execute('INSERT INTO workspace VALUES (?,?,?) ON CONFLICT(account) DO UPDATE SET revision=excluded.revision,state=excluded.state', (account, revision, text))
        return {'revision': revision, 'mode': 'preview'}

    def report(self, account: str, expected_revision: int, scope: str) -> dict:
        current = self.get(account)
        if current['revision'] != expected_revision:
            raise RevisionConflict('Workspace changed. Read it again before generating a report.')
        state = current['state']
        if state is None:
            raise ValueError('Open the dashboard once to initialise the sample workspace.')
        orders = [o for o in state['orders'] if scope == 'All sample dates' or o['date'] == scope]
        report = {'id': str(uuid4()), 'date': scope if scope != 'All sample dates' else datetime.now(timezone(timedelta(hours=8))).date().isoformat(), 'range': scope,
                  'orders': orders, 'attention': [f"{t['customer']}: {t['reason']}" for t in state['threads'] if not t['resolved']],
                  'coverage': f'Fictional UI dataset only. {len(orders)} records in {scope}. No live WhatsApp history analysed.',
                  'created': datetime.now(timezone.utc).isoformat()}
        data = ReportData.model_validate(report)
        content = build_report(data)
        state['reports'].insert(0, report)
        self.save(account, expected_revision, state)
        return {'report': report, 'content': content}


def build_report(report: ReportData) -> bytes:
    """Create and reopen an actual XLSX; totals include only confirmed orders."""
    if len({o.id for o in report.orders}) != len(report.orders):
        raise ValueError('Duplicate order IDs; repair the report before exporting.')
    if any(o.status == 'Confirmed' and not o.delivery.strip() for o in report.orders):
        raise ValueError('Confirmed orders require delivery details.')
    wb = Workbook()
    orders = wb.active
    orders.title = 'Orders'
    orders.append(['Sample data — not live WhatsApp sales'])
    orders.append(['Order', 'Customer', 'Item', 'Quantity', 'Unit price MYR', 'Total MYR', 'Status', 'Payment', 'Delivery', 'Date', 'Source', 'Revision'])
    confirmed = Decimal(0)
    verified = Decimal(0)
    for o in report.orders:
        amount = (o.quantity * o.unitPrice).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        if o.status == 'Confirmed':
            confirmed += amount
            if o.payment == 'Verified': verified += amount
        orders.append([o.id,o.customer,o.item,o.quantity,float(o.unitPrice),float(amount),o.status,o.payment,o.delivery,o.date,o.source,o.revision])
    summary = wb.create_sheet('Summary')
    summary.append(['Metric','Value'])
    for row in [('Dataset','Fictional UI preview'),('Range',report.range),('Timezone','Asia/Singapore'),('Records',len(report.orders)),('Confirmed orders',sum(o.status=='Confirmed' for o in report.orders)),('Confirmed value MYR',float(confirmed)),('Verified payments MYR',float(verified)),('Coverage',report.coverage),('Created',report.created)]: summary.append(row)
    attention = wb.create_sheet('Needs Attention')
    attention.append(['Unresolved item — sample data'])
    for text in report.attention: attention.append([text[:10000]])
    if not report.attention: attention.append(['No unresolved sample items.'])
    for sheet in wb:
        sheet.freeze_panes = 'A3' if sheet.title == 'Orders' else 'A2'
        for row in sheet:
            for cell in row:
                # Preserve untrusted chat text as text, never an Excel formula.
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        header = 2 if sheet.title == 'Orders' else 1
        for cell in sheet[header]:
            cell.fill = PatternFill('solid', fgColor='254E38')
            cell.font = Font(color='FFFFFF', bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(18, max(len(str(c.value or '')) for c in column[:25]) + 2))
        sheet.auto_filter.ref = f'A{header}:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}'
    for row in orders.iter_rows(min_row=3):
        row[4].number_format = row[5].number_format = '#,##0.00'
    out = BytesIO(); wb.save(out); result = out.getvalue()
    check = load_workbook(BytesIO(result), data_only=True)
    if check.sheetnames != ['Orders','Summary','Needs Attention'] or check['Orders'].max_row != len(report.orders) + 2:
        raise ValueError('Workbook verification failed.')
    if Decimal(str(check['Summary']['B7'].value)) != confirmed or Decimal(str(check['Summary']['B8'].value)) != verified:
        raise ValueError('Workbook totals failed verification.')
    check.close()
    return result
