"""Evidence-bearing workbooks; unknown amounts remain blank, never invented zero."""
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def sheet_rows(analyses):
    rows=[]
    for analysis in analyses:
        for order in analysis['orders']:
            total=None
            if order['quantity'] is not None and order['unit_price'] is not None:
                total=str((Decimal(str(order['unit_price']))*order['quantity']).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
            rows.append({**order,'id':analysis['chat']+'/'+order['anchor'],'chat':analysis['chat'],
                'customer':analysis['customer'],'total':total})
    return rows


def create_workbook(analyses, coverage):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Orders'
    sheet.append(['Order ID','Chat','Customer','Items','Quantity','Unit price','Currency','Total',
                  'Status','Payment claim','Delivery','Missing information','Evidence IDs'])
    evidence = wb.create_sheet('Evidence')
    evidence.append(['Order ID','Chat','Message ID','Exact quote'])
    attention = wb.create_sheet('Needs Attention')
    attention.append(['Chat','Reason'])
    totals, count = {}, 0
    for analysis in analyses:
        for order in analysis['orders']:
            key = analysis['chat'] + '/' + order['anchor']
            amount = None
            if order['quantity'] is not None and order['unit_price'] is not None:
                amount = (Decimal(str(order['unit_price'])) * order['quantity']).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
            sheet.append([key,analysis['chat'],analysis['customer'],order['item'],order['quantity'],order['unit_price'],
                          order['currency'],float(amount) if amount is not None else None,order['status'],order['payment'],
                          order['delivery'],'; '.join(order['missing']),', '.join(e['message_id'] for e in order['evidence'])])
            count += 1
            if order['status'] == 'confirmed' and amount is not None and order['currency']:
                currency = order['currency']
                totals[currency] = totals.get(currency, Decimal(0)) + amount
            for item in order['evidence']:
                evidence.append([key,analysis['chat'],item['message_id'],item['quote']])
            if order['missing']:
                attention.append([analysis['chat'], '; '.join(order['missing'])])
        for limitation in analysis['limitations']:
            attention.append([analysis['chat'], limitation])
    summary = wb.create_sheet('Summary')
    summary.append(['Metric','Value'])
    summary.append(['Records',count])
    summary.append(['Coverage',coverage])
    summary.append(['Payment verification','No payment independently verified; claims only'])
    for currency, total in sorted(totals.items()):
        summary.append(['Confirmed value '+currency, float(total)])
    for ws in wb:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = 's'  # Treat WhatsApp text as text, including = formulas.
        for cell in ws[1]:
            cell.fill = PatternFill('solid',fgColor='254E38')
            cell.font = Font(color='FFFFFF',bold=True)
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 26
    result = BytesIO()
    wb.save(result)
    content = result.getvalue()
    check = load_workbook(BytesIO(content),data_only=True)
    if check['Orders'].max_row != count+1 or check['Summary']['B2'].value != count:
        raise ValueError('Workbook row verification failed.')
    for row in check['Summary'].iter_rows(min_row=5,values_only=True):
        if Decimal(str(row[1])) != totals[row[0].removeprefix('Confirmed value ')]:
            raise ValueError('Workbook total verification failed.')
    return content, count, {key:str(value) for key,value in totals.items()}
